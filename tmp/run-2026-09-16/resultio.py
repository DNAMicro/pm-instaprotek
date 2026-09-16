import openpyxl
F="/home/farsheed/pm-instaprotek/Insta-testing/Regression_QA_Log_PROD-2026-09-16.xlsx"
# columns: 1 Module/Section | 2 Scenario ID | 3 Description | 4 Expected | 5 STATUS | 6 Notes | 7 Defect ID
def load(tab):
    wb=openpyxl.load_workbook(F); ws=wb[tab]; idx={}; sec=None
    for r in range(2, ws.max_row+1):
        a=ws.cell(r,1).value; b=ws.cell(r,2).value
        if a and str(a).strip() and str(a).strip()!='-': sec=str(a).strip()
        if b is None: continue
        try: sid=int(float(str(b).strip()))
        except: continue
        idx[f"{sec}|{sid}"]=r
    return wb, ws, idx

def dump(tab):
    wb, ws, idx = load(tab)
    out=[]
    for k,r in idx.items():
        out.append((k, r, str(ws.cell(r,3).value or '').strip(), str(ws.cell(r,4).value or '').strip()))
    return out

NORM={"PASS":"Pass","FAIL":"Fail","BLOCKED":"Blocked","NA":"N/A","N/A":"N/A",
      "NOT TESTED":"Not Tested","NOT RUN":"Not Tested","PARTIAL":"Pass"}
def write(tab, results, defects=None):
    wb, ws, idx = load(tab); wrote=[]; missed=[]
    for k,(st,note) in results.items():
        r=idx.get(k)
        if not r: missed.append(k); continue
        u=str(st).upper().strip()
        if u=="PARTIAL" and not str(note).startswith("(partial)"): note="(partial) "+str(note)
        ws.cell(r,5).value=NORM.get(u, str(st))
        ws.cell(r,6).value=str(note)[:900]
        ws.cell(r,7).value=(defects or {}).get(k,"None")
        wrote.append(k)
    wb.save(F)
    return len(wrote), missed, idx

def setall(tab, status, note):
    wb, ws, idx = load(tab)
    for k,r in idx.items():
        ws.cell(r,5).value=status; ws.cell(r,6).value=note; ws.cell(r,7).value="None"
    wb.save(F); return len(idx)

def tally(tab):
    wb, ws, idx = load(tab); c={}
    for k,r in idx.items():
        s=str(ws.cell(r,5).value or "Not Tested").strip()
        c[s]=c.get(s,0)+1
    return c, len(idx)
